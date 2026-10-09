"""Собрать catalog.json и JPEG из таблицы аренды/продажи и ZIP с Яндекс Диска.

На компьютере разработчика нужны openpyxl и Pillow:
python tools/build_catalog.py '../Объекты для ТГ (1).xlsx' '/path/to/photos.zip'
Имена выбранных фото зафиксированы после визуальной сверки папок.
"""

import argparse
import json
import re
import zipfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageOps

ROOT = Path(__file__).resolve().parents[1]
BATCH_ID = "catalog-2026-10-08-v2"
CARD_TITLES = {
    2: "Современный офис, пом. 1", 3: "Современный офис, пом. 2",
    4: "Современный офис, пом. 3", 5: "Современный офис, пом. 6",
    6: "Возле метро Сокол", 7: "Апартаменты бизнес-класса рядом с ВДНХ",
    8: "Возле метро Перово", 9: "ЖК Первый Нагатинский",
    10: "Метро Лухмановская", 11: "ЖК Цветочные Поляны Экопарк", 12: "ЖК Алхимово",
}
FOLDERS = {
    2: "Волочаевская д.4/пом 1", 3: "Волочаевская д.4/пом 2",
    4: "Волочаевская д.4/пом 3", 5: "Волочаевская д.4/пом 6",
    6: "Ленинградский п-т д.75 к1Б", 7: "пр-т Мира д.135А",
    8: "3- Владимировская д. 23", 9: "ЖК 1-й Нагатинский",
    10: "8 марта", 11: "Жк Цветочные Поляны", 12: "ЖК Алхимова",
}
# Исключены планы: пом.1 — 29,5 вместо 29,4; пом.2 — 29 вместо 30,5;
# пом.3 — 29 вместо 31; Перово — 36,2 вместо 36,3; Люберцы — 207,3 вместо 214.
# Папка «Запасные фото» в каталоге не используется.
PHOTO_NAMES = {
    2: ["1.png", "2.jpeg", "3.png", "5.png"],
    3: ["1.jpg", "2.jpeg", "3.png", "5.png"],
    4: ["1.jpeg", "2.jpg", "3.png", "5.png"],
    5: ["1.jpg", "2.jpg", "3.png", "4.png", "5.png"],
    6: ["1.JPG", "2.JPG", "3.JPG", "4.JPG", "5.png"],
    7: ["1.jpeg", "2.jpeg", "3.jfif", "4.png", "5.jfif", "6.webp", "7.png"],
    8: ["1.jpg", "2.jpg", "3.png", "4.png", "5.png", " 7.png"],
    9: ["1.webp", "2.webp", "3.webp", "4.webp", "5.webp", "6.png", "7.png"],
    10: ["WhatsApp Image 2026-07-09 at 15.22.14 (1).jpeg",
         "WhatsApp Image 2026-07-10 at 16.38.51.jpeg",
         "WhatsApp Image 2026-07-10 at 16.38.52 (3).jpeg",
         "WhatsApp Image 2026-07-10 at 16.38.52 (4).jpeg",
         "WhatsApp Image 2026-07-10 at 16.38.52.jpeg", "Дизайн без названия (45).png"],
    11: ["2025-01-14_13-34-21.png", "2025-01-14_13-34-39.png",
         "WhatsApp Image 2025-01-14 at 14.11.45 (3).jpeg",
         "WhatsApp Image 2025-01-14 at 14.11.45.jpeg", "image (56).png"],
    12: ["1.png", "2.jpeg", "3.png", "4.png", "5.png", "6.png", "7.png"],
}


def clean_text(value):
    lines = [re.sub(r"[ \t\u00a0]+", " ", line).strip() for line in str(value or "").splitlines()]
    return "\n".join(line for line in lines if line)


def build(xlsx_path, zip_path):
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = workbook.active
    if sheet.cell(1, 5).value != "Стоимость продажи" or sheet.max_row != 12:
        raise ValueError("Ожидается новая таблица из 11 объектов с колонкой «Стоимость продажи»")
    output = ROOT / "catalog.json"
    previous = json.loads(output.read_text(encoding="utf-8")) if output.is_file() else {"items": []}
    old_items = {item["source_row"]: item for item in previous["items"]}
    items = []
    try:
        with zipfile.ZipFile(zip_path) as archive:
            for row in range(2, 13):
                title, address, area, monthly_rent, sale_price, description = (
                    sheet.cell(row, col).value for col in range(1, 7)
                )
                if not all((title, address, area, monthly_rent, sale_price, description)):
                    raise ValueError(f"Не заполнены данные строки {row}")
                folder = "фото объектов/" + FOLDERS[row] + "/"
                selected = [folder + name for name in PHOTO_NAMES[row]]
                photos = []
                for position, name in enumerate(selected, start=1):
                    with Image.open(BytesIO(archive.read(name))) as source:
                        image = ImageOps.exif_transpose(source).convert("RGB")
                        image.thumbnail((1600, 1600), Image.Resampling.LANCZOS)
                        relative = Path("media") / f"object_{row:02d}" / f"{position:02d}.jpg"
                        target = ROOT / relative
                        target.parent.mkdir(parents=True, exist_ok=True)
                        image.save(target, "JPEG", quality=84, optimize=True)
                    if target.stat().st_size > 10_000_000:
                        raise ValueError(f"Фото слишком велико для Telegram: {name}")
                    photos.append(relative.as_posix())
                for stale in target.parent.glob("*.jpg"):
                    if f"media/object_{row:02d}/{stale.name}" not in photos:
                        stale.unlink()
                item = {
                    "source_row": row, "title": CARD_TITLES[row],
                    "title_before_address_cleanup": clean_text(title),
                    "price": clean_text(monthly_rent), "sale_price": clean_text(sale_price),
                    "area": float(area), "location": clean_text(address),
                    "rooms": "Офис" if row <= 5 else ("Апартаменты" if row == 7 else "Коммерческое помещение"),
                    "teaser": "", "description": clean_text(description),
                    "features": [], "photos": photos,
                }
                old = old_items.get(row)
                if old:
                    item["previous_title"] = old.get("previous_title", old["title"])
                    item["previous_location"] = old.get("previous_location", old["location"])
                items.append(item)
    finally:
        workbook.close()
    output.write_text(json.dumps({"batch_id": BATCH_ID, "items": items}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Готово: {len(items)} объектов, {sum(len(item['photos']) for item in items)} фотографий")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("photos_zip", type=Path)
    args = parser.parse_args()
    build(args.xlsx, args.photos_zip)


if __name__ == "__main__":
    main()
